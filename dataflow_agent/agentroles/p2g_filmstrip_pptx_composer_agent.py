"""p2g_filmstrip_pptx_composer_agent
~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~

Film-Strip Pipeline Stage 7: PPTX Composer Agent

将布局信息、PPTX 渲染规格、VLM 图片组装成最终的 PowerPoint 文件。

输入：
- state.layout_json: 布局信息
- state.pptx_render_specs: PPTX 节点结构化规格
- state.vlm_rendered_nodes: VLM 图片路径
- state.node_graph_json: 节点图（包含 edges）
- state.render_plan_json: 渲染计划（包含渲染方式）

输出：
- state.pptx_output_path: 生成的 PPTX 文件路径
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, Optional

from dataflow_agent.state import FilmStripP2GState
from dataflow_agent.logger import get_logger

log = get_logger(__name__)


class FilmStripPPTXBuilder:
    """Film-Strip Pipeline 专用 PPTX 构建器

    与通用 PPTXBuilder 的区别：
    - 从 render_plan_json 获取渲染方式
    - 从 node_graph_json 获取 edges
    """

    def __init__(
        self,
        width_px: int = 1920,
        height_px: int = 1080,
        dpi: int = 96,
    ):
        from pptx import Presentation
        from pptx.util import Emu
        from dataflow_agent.toolkits.pptx_composer.coordinate_utils import CoordinateConverter
        from dataflow_agent.toolkits.pptx_composer.node_renderer import PPTXNodeRenderer
        from dataflow_agent.toolkits.pptx_composer.edge_renderer import EdgeRenderer

        self.width_px = width_px
        self.height_px = height_px
        self.converter = CoordinateConverter(dpi=dpi)

        # 初始化 Presentation
        self.prs = Presentation()

        # 设置幻灯片尺寸
        width_emu, height_emu = self.converter.point_to_emu(width_px, height_px)
        self.prs.slide_width = Emu(width_emu)
        self.prs.slide_height = Emu(height_emu)

        # 添加空白幻灯片
        blank_layout = self.prs.slide_layouts[6]  # 空白布局
        self.slide = self.prs.slides.add_slide(blank_layout)

        # 初始化渲染器
        self.node_renderer = PPTXNodeRenderer(self.slide)
        self.edge_renderer = EdgeRenderer(self.slide)

        # 统计
        self.stats = {
            "pptx_nodes_rendered": 0,
            "vlm_nodes_rendered": 0,
            "edges_rendered": 0,
            "errors": [],
        }

    def build(
        self,
        layout_json: Dict[str, Any],
        render_plan_json: Dict[str, Any],
        pptx_render_specs: Dict[str, Any],
        vlm_rendered_nodes: Dict[str, str],
        node_graph_json: Dict[str, Any],
        render_edges: bool = True,
    ) -> "FilmStripPPTXBuilder":
        """构建 PPTX

        Args:
            layout_json: 布局信息
            render_plan_json: 渲染计划（包含渲染方式）
            pptx_render_specs: PPTX 节点结构化规格
            vlm_rendered_nodes: VLM 图片路径
            node_graph_json: 节点图（包含 edges）
            render_edges: 是否渲染边

        Returns:
            self，支持链式调用
        """
        # 提取节点 bbox 映射
        node_bboxes = self._extract_node_bboxes(layout_json)

        # 提取渲染方式映射
        render_methods = self._extract_render_methods(render_plan_json)

        # Step 1: 渲染边（先渲染边，这样节点会覆盖在边上面）
        if render_edges:
            edges = node_graph_json.get("edges", [])
            if edges:
                count = self.edge_renderer.render_edges(edges, node_bboxes)
                self.stats["edges_rendered"] = count

        # Step 2: 渲染节点
        self._render_all_nodes(
            node_bboxes,
            render_methods,
            pptx_render_specs,
            vlm_rendered_nodes,
        )

        log.info(
            f"PPTX build complete: {self.stats['pptx_nodes_rendered']} pptx nodes, "
            f"{self.stats['vlm_nodes_rendered']} vlm nodes, "
            f"{self.stats['edges_rendered']} edges"
        )

        return self

    def save(self, output_path: str) -> str:
        """保存 PPTX 文件"""
        output_dir = Path(output_path).parent
        output_dir.mkdir(parents=True, exist_ok=True)

        self.prs.save(output_path)
        log.info(f"PPTX saved to: {output_path}")

        return output_path

    def _extract_node_bboxes(self, layout_json: Dict) -> Dict[str, Dict]:
        """提取节点 bbox 映射"""
        bboxes = {}
        positions = layout_json.get("positions", {})

        for node_pos in positions.get("nodes", []):
            node_id = node_pos.get("node_id")
            bbox = node_pos.get("bbox", {})
            if node_id and bbox:
                bboxes[node_id] = bbox

        return bboxes

    def _extract_render_methods(self, render_plan_json: Dict) -> Dict[str, str]:
        """从 render_plan_json 提取渲染方式映射"""
        methods = {}
        by_node = render_plan_json.get("by_node", {})

        for node_id, info in by_node.items():
            methods[node_id] = info.get("render_method", "pptx")

        return methods

    def _render_all_nodes(
        self,
        node_bboxes: Dict[str, Dict],
        render_methods: Dict[str, str],
        pptx_render_specs: Dict[str, Any],
        vlm_rendered_nodes: Dict[str, str],
    ):
        """渲染所有节点"""
        for node_id, bbox in node_bboxes.items():
            method = render_methods.get(node_id, "pptx")
            bbox_emu = self.converter.bbox_to_emu(bbox)

            try:
                if method == "vlm":
                    # VLM 图片
                    image_path = vlm_rendered_nodes.get(node_id)
                    if image_path and Path(image_path).exists():
                        success = self.node_renderer.render_vlm_image(
                            node_id, bbox_emu, image_path,
                            fit_mode="contain"
                        )
                        if success:
                            self.stats["vlm_nodes_rendered"] += 1
                    else:
                        log.warning(f"No VLM image for node {node_id}: {image_path}")
                        self.stats["errors"].append(f"Missing VLM image: {node_id}")

                else:
                    # PPTX 渲染
                    spec = pptx_render_specs.get(node_id)
                    if spec:
                        # 文本框默认启用自适应
                        if spec.get("element_type") == "text_box":
                            spec.setdefault("auto_fit", "shrink")

                        success = self.node_renderer.render_node(
                            node_id, bbox_emu, spec
                        )
                        if success:
                            self.stats["pptx_nodes_rendered"] += 1
                    else:
                        log.warning(f"No PPTX spec for node {node_id}")
                        self.stats["errors"].append(f"Missing PPTX spec: {node_id}")

            except Exception as e:
                log.error(f"Error rendering node {node_id}: {e}")
                self.stats["errors"].append(f"Render error {node_id}: {str(e)}")


async def p2g_filmstrip_pptx_composer_agent(
    state: FilmStripP2GState,
    output_path: Optional[str] = None,
    render_edges: bool = True,
    **kwargs,
) -> FilmStripP2GState:
    """Film-Strip PPTX Composer Agent 入口函数

    纯计算，不调用 LLM。

    Args:
        state: FilmStripP2GState 状态对象
        output_path: 输出路径，默认自动生成
        render_edges: 是否渲染连接线

    Returns:
        更新后的 FilmStripP2GState
    """
    # 获取输入
    layout_json = getattr(state, "layout_json", {})
    render_plan_json = getattr(state, "render_plan_json", {})
    pptx_render_specs = getattr(state, "pptx_render_specs", {})
    vlm_rendered_nodes = getattr(state, "vlm_rendered_nodes", {})
    node_graph_json = getattr(state, "node_graph_json", {})

    canvas_width = getattr(state.request, "canvas_width", 1920)
    canvas_height = getattr(state.request, "canvas_height", 1080)
    output_dir = getattr(state.request, "output_dir", ".tmp/filmstrip_p2g")

    # 验证输入
    if not layout_json or not layout_json.get("positions", {}).get("nodes"):
        log.error("[PPTXComposer] layout_json is empty or has no nodes")
        state.agent_results["p2g_filmstrip_pptx_composer_agent"] = {
            "status": "failed",
            "error": "layout_json is empty or has no nodes",
        }
        return state

    # 默认输出路径
    if not output_path:
        output_path = str(Path(output_dir) / "filmstrip_output.pptx")

    log.info("[PPTXComposer] Starting PPTX composition...")
    log.info(f"  - Output path: {output_path}")
    log.info(f"  - Slide size: {canvas_width}x{canvas_height}")
    log.info(f"  - Render edges: {render_edges}")
    log.info(f"  - VLM nodes: {len(vlm_rendered_nodes)}")
    log.info(f"  - PPTX specs: {len(pptx_render_specs)}")

    try:
        # 构建 PPTX
        builder = FilmStripPPTXBuilder(
            width_px=canvas_width,
            height_px=canvas_height,
        )

        builder.build(
            layout_json=layout_json,
            render_plan_json=render_plan_json,
            pptx_render_specs=pptx_render_specs,
            vlm_rendered_nodes=vlm_rendered_nodes,
            node_graph_json=node_graph_json,
            render_edges=render_edges,
        )

        # 保存
        saved_path = builder.save(output_path)

        # 更新 state
        setattr(state, "pptx_output_path", saved_path)
        state.agent_results["p2g_filmstrip_pptx_composer_agent"] = {
            "status": "ok",
            "stats": builder.stats,
            "output_path": saved_path,
        }

        log.info(f"[PPTXComposer] PPTX composition complete: {saved_path}")

    except Exception as e:
        log.error(f"[PPTXComposer] PPTX composition failed: {e}")
        state.agent_results["p2g_filmstrip_pptx_composer_agent"] = {
            "status": "failed",
            "error": str(e),
        }
        raise

    return state
