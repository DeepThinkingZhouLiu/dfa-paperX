"""p2g_chunk_pptx_composer_agent
~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~

Chunk 级别的 PPTX 组装器。

将 chunk 图片按照 bbox 位置放置到 PPTX 幻灯片中。

支持两种 State:
- Paper2GraphState: 旧版 node-based pipeline
- ChunkP2GState: 新版 chunk-based pipeline

输入：
- state.chunk_images: {chunk_id: {path, bbox}} 字典
- state.chunk_layout_json 或 state.layout_json: 包含 canvas 尺寸信息

输出：
- state.pptx_output_path: 生成的 PPTX 文件路径
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any, Dict, Optional

from pptx import Presentation
from pptx.util import Emu

from dataflow_agent.state import MainState, Paper2GraphState, ChunkP2GState
from dataflow_agent.toolkits.tool_manager import ToolManager
from dataflow_agent.logger import get_logger
from dataflow_agent.agentroles.base_agent import BaseAgent
from dataflow_agent.agentroles.registry import register

log = get_logger(__name__)

# 默认输出路径
DEFAULT_OUTPUT_PATH = ".tmp/output/paper2graph_chunk.pptx"

# 默认 DPI
DEFAULT_DPI = 96


class ChunkPPTXBuilder:
    """Chunk 级别的 PPTX 构建器

    将 chunk 图片按照 bbox 位置放置到幻灯片中。
    """

    def __init__(
        self,
        width_px: int = 1920,
        height_px: int = 1080,
        dpi: int = DEFAULT_DPI,
    ):
        """
        Args:
            width_px: 幻灯片宽度（像素）
            height_px: 幻灯片高度（像素）
            dpi: DPI 设置
        """
        self.width_px = width_px
        self.height_px = height_px
        self.dpi = dpi

        # 初始化 Presentation
        self.prs = Presentation()

        # 设置幻灯片尺寸
        width_emu = self._px_to_emu(width_px)
        height_emu = self._px_to_emu(height_px)
        self.prs.slide_width = Emu(width_emu)
        self.prs.slide_height = Emu(height_emu)

        # 添加空白幻灯片
        blank_layout = self.prs.slide_layouts[6]  # 空白布局
        self.slide = self.prs.slides.add_slide(blank_layout)

        # 统计
        self.stats = {
            "chunks_rendered": 0,
            "errors": [],
        }

    def _px_to_emu(self, px: int) -> int:
        """像素转 EMU"""
        return int(px * 914400 / self.dpi)

    def add_chunk_image(
        self,
        chunk_id: str,
        image_path: str,
        bbox: Dict[str, int],
    ) -> bool:
        """添加 chunk 图片到幻灯片

        Args:
            chunk_id: chunk ID
            image_path: 图片路径
            bbox: {x, y, w, h} 像素坐标

        Returns:
            是否添加成功
        """
        if not image_path or not Path(image_path).exists():
            log.error(f"Image not found for chunk {chunk_id}: {image_path}")
            self.stats["errors"].append(f"Missing image: {chunk_id}")
            return False

        try:
            # 转换坐标
            left_emu = self._px_to_emu(bbox.get("x", 0))
            top_emu = self._px_to_emu(bbox.get("y", 0))
            width_emu = self._px_to_emu(bbox.get("w", 400))
            height_emu = self._px_to_emu(bbox.get("h", 300))

            # 添加图片
            self.slide.shapes.add_picture(
                image_path,
                Emu(left_emu),
                Emu(top_emu),
                Emu(width_emu),
                Emu(height_emu),
            )

            log.debug(f"Added chunk image: {chunk_id} at ({bbox['x']}, {bbox['y']}) size {bbox['w']}x{bbox['h']}")
            self.stats["chunks_rendered"] += 1
            return True

        except Exception as e:
            log.error(f"Failed to add chunk image {chunk_id}: {e}")
            self.stats["errors"].append(f"Render error {chunk_id}: {str(e)}")
            return False

    def save(self, output_path: str) -> str:
        """保存 PPTX 文件

        Args:
            output_path: 输出路径

        Returns:
            实际保存的路径
        """
        # 确保目录存在
        output_dir = Path(output_path).parent
        output_dir.mkdir(parents=True, exist_ok=True)

        self.prs.save(output_path)
        log.info(f"PPTX saved to: {output_path}")

        return output_path


@register("p2g_chunk_pptx_composer_agent")
class P2gChunkPptxComposerAgent(BaseAgent):
    """Chunk 级别 PPTX 组装器 Agent

    将 chunk 图片按照 bbox 位置放置到 PPTX 幻灯片中。
    """

    def __init__(self, tool_manager: Optional[ToolManager] = None, **kwargs):
        super().__init__(tool_manager=tool_manager, **kwargs)

    @classmethod
    def create(cls, tool_manager: Optional[ToolManager] = None, **kwargs):
        return cls(tool_manager=tool_manager, **kwargs)

    @property
    def role_name(self) -> str:
        return "p2g_chunk_pptx_composer_agent"

    @property
    def system_prompt_template_name(self) -> str:
        return ""  # 不使用 LLM

    @property
    def task_prompt_template_name(self) -> str:
        return ""  # 不使用 LLM

    def get_default_pre_tool_results(self) -> Dict[str, Any]:
        return {}

    def get_task_prompt_params(self, pre_tool_results: Dict[str, Any]) -> Dict[str, Any]:
        return {}

    def update_state_result(
        self,
        state: MainState,
        result: Dict[str, Any],
        pre_tool_results: Dict[str, Any],
    ):
        """更新 state.pptx_output_path"""
        if isinstance(result, dict):
            if "output_path" in result:
                setattr(state, "pptx_output_path", result["output_path"])
        state.agent_results[self.role_name] = result


async def p2g_chunk_pptx_composer_agent(
    state: Paper2GraphState,
    output_path: Optional[str] = None,
    slide_width: int = 1920,
    slide_height: int = 1080,
) -> Paper2GraphState:
    """Chunk PPTX 组装器入口函数

    从 state.chunk_images 中获取每个 chunk 的图片和 bbox，
    将它们放置到 PPTX 幻灯片中。

    Args:
        state: Paper2GraphState 状态对象
        output_path: 输出路径，默认 .tmp/output/paper2graph_chunk.pptx
        slide_width: 幻灯片宽度（像素）
        slide_height: 幻灯片高度（像素）

    Returns:
        更新后的 Paper2GraphState
    """
    # 获取 chunk_images
    chunk_images = getattr(state, "chunk_images", {}) or {}

    if not chunk_images:
        log.warning("chunk_images is empty, cannot compose PPTX")
        state.agent_results["p2g_chunk_pptx_composer_agent"] = {
            "status": "failed",
            "error": "No chunk images to compose",
        }
        return state

    # 从 layout_json 获取 canvas 尺寸
    layout_json = getattr(state, "layout_json", {}) or {}
    canvas = layout_json.get("canvas", {})
    if canvas:
        slide_width = canvas.get("width", slide_width)
        slide_height = canvas.get("height", slide_height)

    # 设置输出路径
    if output_path is None:
        output_path = DEFAULT_OUTPUT_PATH

    log.info(f"[Chunk PPTX Composer] Starting composition...")
    log.info(f"[Chunk PPTX Composer] Output path: {output_path}")
    log.info(f"[Chunk PPTX Composer] Slide size: {slide_width}x{slide_height}")
    log.info(f"[Chunk PPTX Composer] Chunks to compose: {len(chunk_images)}")

    try:
        # 构建 PPTX
        builder = ChunkPPTXBuilder(
            width_px=slide_width,
            height_px=slide_height,
        )

        # 添加所有 chunk 图片
        for chunk_id, chunk_data in chunk_images.items():
            image_path = chunk_data.get("path", "")
            bbox = chunk_data.get("bbox", {})

            if not bbox:
                log.warning(f"No bbox for chunk {chunk_id}, skipping")
                continue

            builder.add_chunk_image(chunk_id, image_path, bbox)

        # 保存
        saved_path = builder.save(output_path)

        # 更新 state
        setattr(state, "pptx_output_path", saved_path)

        state.agent_results["p2g_chunk_pptx_composer_agent"] = {
            "status": "ok",
            "output_path": saved_path,
            "stats": builder.stats,
        }

        log.info(f"[Chunk PPTX Composer] Completed: {builder.stats['chunks_rendered']} chunks rendered")

    except Exception as e:
        log.error(f"[Chunk PPTX Composer] Failed: {e}")
        state.agent_results["p2g_chunk_pptx_composer_agent"] = {
            "status": "failed",
            "error": str(e),
        }
        raise

    return state


# ============================================================================
# ChunkP2GState 专用入口函数
# ============================================================================

async def p2g_chunk_pptx_composer_for_chunk_pipeline(
    state: ChunkP2GState,
    output_path: Optional[str] = None,
) -> ChunkP2GState:
    """ChunkP2GState 专用的 PPTX 组装入口函数

    Args:
        state: ChunkP2GState 状态对象
        output_path: 输出路径

    Returns:
        更新后的 ChunkP2GState
    """
    chunk_images = getattr(state, "chunk_images", {}) or {}
    chunk_layout_json = getattr(state, "chunk_layout_json", {}) or {}

    if not chunk_images:
        log.warning("[ChunkPptxComposer] chunk_images is empty")
        state.agent_results["p2g_chunk_pptx_composer_agent"] = {
            "status": "failed",
            "error": "No chunk images to compose",
        }
        return state

    # 获取 canvas 尺寸
    canvas = chunk_layout_json.get("canvas", {})
    slide_width = canvas.get("width", 1920)
    slide_height = canvas.get("height", 1080)

    # 设置输出路径
    if output_path is None:
        output_dir = getattr(state.request, "output_dir", ".tmp/chunk_pipeline_output")
        output_path = os.path.join(output_dir, "paper2graph_chunk.pptx")

    log.info(f"[ChunkPptxComposer] Starting: {len(chunk_images)} chunks, size {slide_width}x{slide_height}")

    try:
        # 构建 PPTX
        builder = ChunkPPTXBuilder(width_px=slide_width, height_px=slide_height)

        # 添加所有 chunk 图片
        chunks_layout = {c["chunk_id"]: c for c in chunk_layout_json.get("chunks", [])}

        for chunk_id, chunk_data in chunk_images.items():
            image_path = chunk_data.get("path", "")
            # 从 chunk_layout_json 获取 bbox
            layout = chunks_layout.get(chunk_id, {})
            bbox = layout.get("bbox", {})

            if bbox:
                builder.add_chunk_image(chunk_id, image_path, bbox)
            else:
                log.warning(f"[ChunkPptxComposer] No bbox for {chunk_id}, skipping")

        # 保存
        saved_path = builder.save(output_path)

        # 更新 state
        setattr(state, "pptx_output_path", saved_path)
        setattr(state, "extracted_elements", {})  # MVP 版本暂不提取

        state.agent_results["p2g_chunk_pptx_composer_agent"] = {
            "status": "ok",
            "output_path": saved_path,
            "stats": builder.stats,
        }

        log.info(f"[ChunkPptxComposer] Completed: {builder.stats['chunks_rendered']} chunks")

    except Exception as e:
        log.error(f"[ChunkPptxComposer] Failed: {e}")
        state.agent_results["p2g_chunk_pptx_composer_agent"] = {
            "status": "failed",
            "error": str(e),
        }

    return state
