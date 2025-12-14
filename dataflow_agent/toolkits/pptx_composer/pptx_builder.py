"""PPTX 构建器

核心类，负责组装所有元素生成最终的 PowerPoint 文件。
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any, Dict, List, Optional

from pptx import Presentation
from pptx.util import Emu

from dataflow_agent.logger import get_logger
from .coordinate_utils import CoordinateConverter
from .node_renderer import PPTXNodeRenderer
from .edge_renderer import EdgeRenderer

log = get_logger(__name__)


class PPTXBuilder:
    """PPTX 文档构建器
    
    组装布局、节点、边生成最终的 PowerPoint 文件。
    """
    
    def __init__(
        self,
        width_px: int = 1920,
        height_px: int = 1080,
        dpi: int = 96,
    ):
        """
        Args:
            width_px: 幻灯片宽度（像素）
            height_px: 幻灯片高度（像素）
            dpi: DPI 设置
        """
        self.width_px = width_px
        self.height_px = height_px
        self.converter = CoordinateConverter(dpi=dpi)
        
        # 初始化 Presentation
        self.prs = Presentation()
        
        # 设置幻灯片尺寸
        width_emu, height_emu = self.converter.point_to_emu(width_px, height_px)
        self.prs.slide_width = Emu(width_emu)
        self.prs.slide_height = Emu(height_emu)
        
        # 添加空白幻灯片
        blank_layout = self.prs.slide_layouts[6]  # 空白布局
        self.slide = self.prs.slides.add_slide(blank_layout)
        
        # 初始化渲染器
        self.node_renderer = PPTXNodeRenderer(self.slide)
        self.edge_renderer = EdgeRenderer(self.slide)
        
        # 统计
        self.stats = {
            "pptx_nodes_rendered": 0,
            "vlm_nodes_rendered": 0,
            "edges_rendered": 0,
            "errors": [],
        }
    
    def build(
        self,
        layout_json: Dict[str, Any],
        node_render_design: Dict[str, Any],
        pptx_render_specs: Dict[str, Any],
        vlm_rendered_nodes: Dict[str, str],
        semantic_json: Dict[str, Any],
        render_edges: bool = True,
        render_chunk_bg: bool = False,
    ) -> "PPTXBuilder":
        """构建 PPTX
        
        Args:
            layout_json: 布局信息
            node_render_design: 渲染设计
            pptx_render_specs: PPTX 节点结构化规格
            vlm_rendered_nodes: VLM 图片路径
            semantic_json: 语义信息（包含 edges）
            render_edges: 是否渲染边
            render_chunk_bg: 是否渲染 chunk 背景
            
        Returns:
            self，支持链式调用
        """
        # 提取节点 bbox 映射
        node_bboxes = self._extract_node_bboxes(layout_json)
        
        # 提取渲染方式映射
        render_methods = self._extract_render_methods(node_render_design)
        
        # Step 1: 渲染 chunk 背景（可选）
        if render_chunk_bg:
            self._render_chunk_backgrounds(layout_json)
        
        # Step 2: 渲染边
        if render_edges:
            edges = semantic_json.get("edges", [])
            if edges:
                count = self.edge_renderer.render_edges(edges, node_bboxes)
                self.stats["edges_rendered"] = count
        
        # Step 3: 渲染节点
        self._render_all_nodes(
            node_bboxes,
            render_methods,
            pptx_render_specs,
            vlm_rendered_nodes,
        )
        
        log.info(
            f"PPTX build complete: {self.stats['pptx_nodes_rendered']} pptx nodes, "
            f"{self.stats['vlm_nodes_rendered']} vlm nodes, "
            f"{self.stats['edges_rendered']} edges"
        )
        
        return self
    
    def save(self, output_path: str) -> str:
        """保存 PPTX 文件
        
        Args:
            output_path: 输出路径
            
        Returns:
            实际保存的路径
        """
        # 确保目录存在
        output_dir = Path(output_path).parent
        output_dir.mkdir(parents=True, exist_ok=True)
        
        self.prs.save(output_path)
        log.info(f"PPTX saved to: {output_path}")
        
        return output_path
    
    def _extract_node_bboxes(self, layout_json: Dict) -> Dict[str, Dict]:
        """提取节点 bbox 映射"""
        bboxes = {}
        positions = layout_json.get("positions", {})
        
        for node_pos in positions.get("nodes", []):
            node_id = node_pos.get("node_id")
            bbox = node_pos.get("bbox", {})
            if node_id and bbox:
                bboxes[node_id] = bbox
        
        return bboxes
    
    def _extract_render_methods(self, node_render_design: Dict) -> Dict[str, str]:
        """提取节点渲染方式映射"""
        methods = {}

        for chunk in node_render_design.get("chunks", []):
            for node in chunk.get("nodes", []):
                node_id = node.get("node_id")
                method = node.get("render_method", "pptx")
                if node_id:
                    methods[node_id] = method

        return methods

    def _render_chunk_backgrounds(self, layout_json: Dict):
        """渲染 chunk 背景（可选功能）"""
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.dml.color import RGBColor

        positions = layout_json.get("positions", {})

        for chunk_pos in positions.get("chunks", []):
            chunk_id = chunk_pos.get("chunk_id")
            bbox = chunk_pos.get("bbox", {})

            if not bbox:
                continue

            bbox_emu = self.converter.bbox_to_emu(bbox)
            left, top, width, height = bbox_emu

            try:
                shape = self.slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE,
                    Emu(left), Emu(top), Emu(width), Emu(height)
                )

                # 淡色背景
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor(0xF5, 0xF5, 0xF5)
                shape.line.fill.background()  # 无边框

                log.debug(f"Rendered chunk background: {chunk_id}")

            except Exception as e:
                log.warning(f"Failed to render chunk background {chunk_id}: {e}")

    def _render_all_nodes(
        self,
        node_bboxes: Dict[str, Dict],
        render_methods: Dict[str, str],
        pptx_render_specs: Dict[str, Any],
        vlm_rendered_nodes: Dict[str, str],
    ):
        """渲染所有节点"""
        for node_id, bbox in node_bboxes.items():
            method = render_methods.get(node_id, "pptx")
            bbox_emu = self.converter.bbox_to_emu(bbox)

            try:
                if method == "vlm":
                    # VLM 图片 - 使用 contain 模式，等比缩放+居中
                    image_path = vlm_rendered_nodes.get(node_id)
                    if image_path:
                        success = self.node_renderer.render_vlm_image(
                            node_id, bbox_emu, image_path,
                            fit_mode="contain"
                        )
                        if success:
                            self.stats["vlm_nodes_rendered"] += 1
                    else:
                        log.warning(f"No VLM image for node {node_id}")
                        self.stats["errors"].append(f"Missing VLM image: {node_id}")

                else:
                    # PPTX 渲染
                    spec = pptx_render_specs.get(node_id)
                    if spec:
                        # 文本框默认启用自适应
                        if spec.get("element_type") == "text_box":
                            spec.setdefault("auto_fit", "shrink")

                        success = self.node_renderer.render_node(
                            node_id, bbox_emu, spec
                        )
                        if success:
                            self.stats["pptx_nodes_rendered"] += 1
                    else:
                        log.warning(f"No PPTX spec for node {node_id}")
                        self.stats["errors"].append(f"Missing PPTX spec: {node_id}")

            except Exception as e:
                log.error(f"Error rendering node {node_id}: {e}")
                self.stats["errors"].append(f"Render error {node_id}: {str(e)}")
